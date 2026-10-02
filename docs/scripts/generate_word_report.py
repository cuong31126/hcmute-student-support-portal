import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(165, 28, 48) # Harvard Crimson
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(17, 24, 39) # Deep Slate
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(55, 65, 81)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    return h

def add_paragraph_styled(doc, text="", bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(17, 24, 39)
    if text:
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(31, 41, 55)
    return p

def add_image_placeholder(doc, fig_number, fig_title, puml_file, instructions="Dán hình ảnh xuất ra từ PlantUML hoặc công cụ vẽ vào khung này"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=200, bottom=200, left=300, right=300)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="12" w:space="0" w:color="9CA3AF"/>'
        f'<w:left w:val="dashed" w:sz="12" w:space="0" w:color="9CA3AF"/>'
        f'<w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="9CA3AF"/>'
        f'<w:right w:val="dashed" w:sz="12" w:space="0" w:color="9CA3AF"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    
    r1 = p.add_run("🖼️ [CHỖ DÁN HÌNH ẢNH MINH HỌA]")
    r1.font.name = 'Arial'
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(165, 28, 48)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(f"{fig_number}: {fig_title}")
    r2.font.name = 'Arial'
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(17, 24, 39)
    
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(2)
    r3 = p3.add_run(f"📁 Tệp mã nguồn UML tương ứng: docs/uml/{puml_file}")
    r3.font.name = 'Arial'
    r3.font.size = Pt(10)
    r3.font.italic = True
    r3.font.color.rgb = RGBColor(37, 99, 235)
    
    p4 = cell.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(2)
    p4.paragraph_format.space_after = Pt(8)
    r4 = p4.add_run(f"👉 Hướng dẫn thao tác: {instructions}")
    r4.font.name = 'Arial'
    r4.font.size = Pt(9.5)
    r4.font.color.rgb = RGBColor(107, 114, 128)
    
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(14)
    r_cap = caption.add_run(f"{fig_number}. {fig_title}")
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(55, 65, 81)

def add_table_data(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, "A51C30") # Crimson header
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Body rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(31, 41, 55)
            
    # Set borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="D1D5DB"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="9CA3AF"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    return table

def build_word_document():
    doc = docx.Document()
    
    # Set margins 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title Cover
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_uni.font.name = 'Arial'
    r_uni.font.size = Pt(12)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(55, 65, 81)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BÁO CÁO KHẢO SÁT HIỆN TRẠNG, ĐẶC TẢ YÊU CẦU\nVÀ THIẾT KẾ HỆ THỐNG")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(165, 28, 48)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("ĐỀ TÀI: CỔNG TƯ VẤN HỌC VỤ SINH VIÊN, QUẢN LÝ TICKET SLA\n& TRỢ LÝ AI RAG (QAUTE PORTAL)")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(17, 24, 39)
    
    doc.add_paragraph()
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("Nền tảng: Spring Boot 3.3, MySQL 8, Python FastAPI ChromaDB, Thymeleaf, Bootstrap 5\nPhiên bản tài liệu: 2.0-FINAL | Năm học: 2026")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(11)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(107, 114, 128)
    
    doc.add_page_break()
    
    # CHƯƠNG 2
    add_heading_styled(doc, "CHƯƠNG 2: KHẢO SÁT HIỆN TRẠNG & ĐÁNH GIÁ NHU CẦU THỰC TẾ", level=1)
    
    add_heading_styled(doc, "2.1. Khảo sát hiện trạng hỗ trợ học vụ tại HCMUTE", level=2)
    add_paragraph_styled(doc, "Trường Đại học Sư phạm Kỹ thuật TP. Hồ Chí Minh (HCMUTE) là cơ sở giáo dục đại học công lập trọng điểm với quy mô hơn 25.000 sinh viên cùng hàng chục ngàn thí sinh, phụ huynh quan tâm mỗi kỳ tuyển sinh. Hiện nay, công tác truyền thông, giải đáp thắc mắc và hỗ trợ thủ tục học vụ được phân bổ qua nhiều đơn vị đầu mối trực thuộc:")
    add_paragraph_styled(doc, "Đầu mối thông tin về đề án tuyển sinh các hệ đào tạo, học phí, điểm chuẩn xét tuyển, giải đáp thắc mắc của thí sinh và phụ huynh qua Hotline, Fanpage Facebook và các buổi tư vấn trực tiếp.", bold_prefix="1. Phòng Tuyển sinh và Truyền thông (Phòng A1-101): ")
    add_paragraph_styled(doc, "Tiếp nhận và giải quyết đăng ký môn học, lịch thi, hoãn thi, phúc khảo điểm thi, chứng chỉ ngoại ngữ chuẩn đầu ra (TOEIC, IELTS), học bổng khuyến khích học tập, trợ cấp xã hội và xét tốt nghiệp.", bold_prefix="2. Phòng Đào tạo & Phòng Công tác Sinh viên (Phòng A1-201): ")
    add_paragraph_styled(doc, "Tiếp nhận thông tin phong trào, hoạt động tình nguyện, rèn luyện kỹ năng mềm, xác nhận điểm rèn luyện (ĐRL) và các chương trình truyền thông đa phương tiện.", bold_prefix="3. Đoàn Thanh niên - Hội Sinh viên trường (Phòng A1-102): ")
    add_paragraph_styled(doc, "Hướng dẫn đồ án môn học, khóa luận tốt nghiệp, giới thiệu thực tập doanh nghiệp và giải quyết các vướng mắc chuyên ngành.", bold_prefix="4. Các Khoa chuyên môn (Khoa CNTT - Tòa E1, Cơ khí, Điện - Điện tử, Ngoại ngữ...): ")
    
    add_image_placeholder(doc, "Hình 2-1", "Giao diện các kênh hỗ trợ học vụ & Fanpage HCMUTE hiện tại", "placeholder_hcmute.puml", "Chụp ảnh màn hình trang tuyển sinh, cổng đào tạo và fanpage HCMUTE hiện tại rồi dán vào đây")
    
    add_heading_styled(doc, "2.2. Khảo sát các giải pháp cổng thông tin học vụ trong và ngoài nước", level=2)
    add_paragraph_styled(doc, "Nhóm nghiên cứu đã khảo sát và đánh giá 3 nhóm mô hình hỗ trợ sinh viên phổ biến:")
    add_paragraph_styled(doc, "Triển khai mô hình One-Stop Service tích hợp vào cổng đào tạo. Cho phép sinh viên nộp đơn trực tuyến đối với các nghiệp vụ hành chính cơ bản.", bold_prefix="1. Cổng thông tin một cửa - ĐHQG TP.HCM & HCMUT (MyBK): ")
    add_paragraph_styled(doc, "Ứng dụng kênh tiếp nhận Ticket (Helpdesk) kết hợp với chatbot rule-based (dựa trên kịch bản nút bấm cố định). Hạn chế lớn là chatbot không hiểu được tiếng Việt tự nhiên và không liên thông diễn đàn.", bold_prefix="2. Hệ thống Hỗ trợ sinh viên - Trường ĐH FPT & HUTECH: ")
    add_paragraph_styled(doc, "Hệ thống quản lý yêu cầu tiêu chuẩn ITIL với SLA đa tầng, phân quyền theo nhóm hỗ trợ. Tuy nhiên, chi phí bản quyền quá cao và thiếu khả năng tùy biến sâu cho quy trình học vụ và AI RAG tiếng Việt.", bold_prefix="3. Các nền tảng Service Desk tiêu chuẩn doanh nghiệp (Zendesk, Jira Service Management): ")
    
    add_image_placeholder(doc, "Hình 2-2", "Mô hình One-Stop Student Portal và Hệ thống Service Desk tiêu chuẩn", "placeholder_service_desk.puml", "Dán ảnh sơ đồ đối sánh mô hình Service Desk học vụ vào đây")

    add_heading_styled(doc, "2.3. Bảng so sánh đánh giá ưu - nhược điểm các hệ thống hiện nay", level=2)
    table_comp_headers = ["Tiêu chí so sánh", "Kênh HCMUTE hiện tại", "One-Stop ĐH lớn (Bách Khoa/FPT)", "Helpdesk Doanh nghiệp (Zendesk)", "Dự án QAUTE Portal đề xuất"]
    table_comp_rows = [
        ["Kênh tiếp nhận", "Phân tán (Email, Fanpage, Trực tiếp)", "Tập trung trên Web Portal", "Tập trung qua Portal / Email", "Tập trung đa kênh: Form Ticket + Chat AI"],
        ["Cam kết thời hạn (SLA)", "Không có cam kết hạn chót", "Thời gian ước tính (không tự động)", "Tự động tính hạn chót, đếm ngược SLA", "Tự động gán SLA theo mức ưu tiên (24h/72h/7d)"],
        ["Trợ lý AI hỏi đáp", "Không có hoặc bot nút bấm cứng nhắc", "Chatbot kịch bản tĩnh", "AI trả lời theo bài mẫu sẵn", "AI RAG kép (Dual-Engine) đọc hiểu PDF quy chế"],
        ["Chống quá tải & Cản lọc", "100% cán bộ phải đọc và trả lời", "Cản lọc được khoảng 20-30% câu cơ bản", "Phụ thuộc vào Knowledge Base tĩnh", "Phễu cản tải 4 tầng (Cache+FAQ+RAG), giảm >= 70%"],
        ["Bảng tin & Diễn đàn", "Tách rời, không có diễn đàn sinh viên", "Có thông báo, không có diễn đàn", "Không hỗ trợ diễn đàn cộng đồng", "2 luồng độc lập: Bảng tin Cán bộ & Diễn đàn duyệt"],
        ["Đa phương tiện & Video", "Nhúng link thủ công", "Tải file đính kèm đơn giản", "Lưu file đính kèm cơ bản", "Tự động nhận Video MP4 từ Webhook Node.js"],
        ["Phân quyền Khoa/Phòng", "Thủ công trực tiếp tại phòng", "Phân quyền theo chức năng quản trị", "Phân quyền Queue theo nhóm", "Cách ly dữ liệu nghiêm ngặt theo department_id"]
    ]
    add_table_data(doc, table_comp_headers, table_comp_rows)
    
    add_heading_styled(doc, "2.4. Xác định bài toán cốt lõi & Tính cấp thiết của dự án QAUTE Portal", level=2)
    add_paragraph_styled(doc, "Từ kết quả khảo sát thực trạng, việc phát triển QAUTE Portal giải quyết trọn vẹn 4 bài toán then chốt:")
    add_paragraph_styled(doc, "Giải phóng sức lao động của cán bộ tư vấn bằng Trợ lý AI RAG có khả năng đọc hiểu ngữ nghĩa từ hàng trăm trang công văn, quy chế học bổng, chuyển đổi điểm ngoại ngữ.", bold_prefix="1. Tự động hóa giải đáp học vụ 24/7: ")
    add_paragraph_styled(doc, "Đảm bảo mọi thắc mắc của sinh viên/thí sinh đều được cấp mã tra cứu, định tuyến chính xác về Khoa/Phòng phụ trách và cam kết thời hạn giải quyết đúng hạn.", bold_prefix="2. Chuẩn hóa quy trình xử lý yêu cầu theo cam kết chất lượng (SLA): ")
    add_paragraph_styled(doc, "Cung cấp bảng tin chính thức có tệp đính kèm và video chuẩn mực, song song với diễn đàn sinh viên có cơ chế tiền kiểm duyệt (Pre-moderation) ngăn ngừa tin tiêu cực.", bold_prefix="3. Môi trường kết nối học đường chính thống và lành mạnh: ")
    add_paragraph_styled(doc, "Kết hợp Java Spring Boot bảo mật cao với Python FastAPI/ChromaDB xử lý vector AI, có khả năng phòng vệ hạn mức API và hoạt động bền bỉ.", bold_prefix="4. Kiến trúc bền vững, an toàn và tối ưu chi phí: ")
    
    doc.add_page_break()
    
    # CHƯƠNG 3
    add_heading_styled(doc, "CHƯƠNG 3: PHÂN TÍCH YÊU CẦU & THIẾT KẾ HỆ THỐNG", level=1)
    
    add_heading_styled(doc, "3.1. Phân tích chức năng theo 4 nhóm Tác nhân (Actors)", level=2)
    
    # Guest
    add_heading_styled(doc, "3.1.1. Phía Khách vãng lai / Thí sinh (GUEST)", level=3)
    guest_headers = ["STT", "Mã chức năng", "Tên chức năng", "Mô tả chi tiết"]
    guest_rows = [
        ["1", "F-GST-01", "Tra cứu thông tin tuyển sinh & Hỏi đáp AI Bot", "Đặt câu hỏi tự nhiên về đề án tuyển sinh, học phí, điểm chuẩn và nhận câu trả lời trích nguồn từ quy chế chính thức."],
        ["2", "F-GST-02", "Gửi Ticket tư vấn qua Email", "Nhập Họ tên, Số điện thoại, Email cá nhân để gửi yêu cầu hỗ trợ mà không cần tài khoản đăng nhập."],
        ["3", "F-GST-03", "Chuyển đổi cuộc hội thoại thành Ticket", "Khi Chatbot AI không thỏa mãn nhu cầu, bấm nút chuyển tiếp toàn bộ ngữ cảnh hội thoại thành Ticket gửi cán bộ."],
        ["4", "F-GST-04", "Tra cứu tiến độ xử lý Ticket qua Token", "Sử dụng đường link bảo mật gửi về Email chứa Token tra cứu để xem tiến trình và phản hồi của cán bộ."],
        ["5", "F-GST-05", "Đánh giá mức độ hài lòng (CSAT)", "Đánh giá chất lượng phục vụ từ 1 đến 5 sao và gửi góp ý sau khi Ticket được giải quyết."],
        ["6", "F-GST-06", "Xem Bảng tin thông báo & Tải tệp công văn", "Xem các thông báo công khai và tải các tệp đính kèm (.pdf, .docx, .xlsx, .mp4)."]
    ]
    add_table_data(doc, guest_headers, guest_rows)
    
    # Student
    add_heading_styled(doc, "3.1.2. Phía Sinh viên chính quy (ROLE_STUDENT)", level=3)
    stu_headers = ["STT", "Mã chức năng", "Tên chức năng", "Mô tả chi tiết"]
    stu_rows = [
        ["1", "F-STU-01", "Đăng ký & Kích hoạt tài khoản bằng Email trường", "Đăng ký tài khoản với email @student.hcmute.edu.vn và kích hoạt bằng mã OTP gửi về hòm thư."],
        ["2", "F-STU-02", "Đăng nhập, Đăng xuất & Quản lý hồ sơ", "Đăng nhập hệ thống, cập nhật thông tin cá nhân, ảnh đại diện, đổi mật khẩu và xem lịch sử tương tác."],
        ["3", "F-STU-03", "Tra cứu Trợ lý AI RAG không giới hạn", "Trò chuyện với Trợ lý AI với hạn mức ưu tiên, hỗ trợ mở rộng từ viết tắt học vụ (ĐRL, ĐKMH, CTĐT...)."],
        ["4", "F-STU-04", "Tạo Ticket hỗ trợ học vụ có đính kèm minh chứng", "Gửi yêu cầu giải quyết vướng mắc (trùng lịch thi, khiếu nại điểm...) kèm tệp đơn từ PDF hoặc hình ảnh chứng minh."],
        ["5", "F-STU-05", "Theo dõi vòng đời Ticket cá nhân", "Quản lý danh sách các Ticket đã gửi, trạng thái hạn chót SLA, trao đổi tin nhắn trực tiếp với Cán bộ."],
        ["6", "F-STU-06", "Đăng bài viết lên Diễn đàn sinh viên", "Soạn bài viết chia sẻ tài liệu, tìm nhóm học tập; bài viết được đưa vào hàng đợi kiểm duyệt (PENDING_APPROVAL)."],
        ["7", "F-STU-07", "Tương tác Thả tim (Like) & Bình luận (Comment)", "Tương tác thả tim, bình luận nhiều cấp trên các bài viết Diễn đàn đã được duyệt công khai."],
        ["8", "F-STU-08", "Báo cáo nội dung vi phạm (Report)", "Báo cáo bài viết hoặc bình luận có nội dung tiêu cực, xuyên tạc hoặc từ ngữ phản cảm lên ban kiểm duyệt."]
    ]
    add_table_data(doc, stu_headers, stu_rows)
    
    # Staff
    add_heading_styled(doc, "3.1.3. Phía Cán bộ Khoa / Phòng ban (ROLE_STAFF)", level=3)
    staff_headers = ["STT", "Mã chức năng", "Tên chức năng", "Mô tả chi tiết"]
    staff_rows = [
        ["1", "F-STF-01", "Đăng nhập phân quyền theo Khoa/Phòng", "Đăng nhập bằng tài khoản Cán bộ gắn mã đơn vị (department_id)."],
        ["2", "F-STF-02", "Dashboard quản lý Ticket theo phạm vi", "Xem danh sách Ticket gửi riêng cho Khoa/Phòng mình, theo dõi nhãn cảnh báo thời hạn SLA (Còn hạn, Sắp quá hạn, Trễ hạn)."],
        ["3", "F-STF-03", "Tiếp nhận xử lý Ticket (Claim Ticket)", "Bấm 'Tiếp nhận' để nhận trách nhiệm xử lý, chuyển trạng thái từ OPEN sang IN_PROGRESS (chống tranh chấp nhận xử lý)."],
        ["4", "F-STF-04", "Trả lời, Hướng dẫn & Đóng Ticket (Resolved)", "Soạn câu trả lời giải đáp, đính kèm văn bản hướng dẫn và chuyển trạng thái sang RESOLVED."],
        ["5", "F-STF-05", "Đăng bài Bảng tin chính thức kèm Video/Tệp", "Soạn thông báo chính thức có đính kèm văn bản và liên kết Video MP4 (tự động nhận qua Webhook từ Node.js)."],
        ["6", "F-STF-06", "Phê duyệt bài viết Diễn đàn sinh viên", "Duyệt (APPROVED) hoặc từ chối kèm lý do (REJECTED) các bài viết sinh viên gửi lên diễn đàn."],
        ["7", "F-STF-07", "Xử lý danh sách báo cáo vi phạm", "Xem các bài viết/bình luận bị tố cáo, thực hiện ẩn bài (HIDDEN) hoặc xóa vi phạm."]
    ]
    add_table_data(doc, staff_headers, staff_rows)
    
    # Admin
    add_heading_styled(doc, "3.1.4. Phía Quản trị viên hệ thống (ROLE_ADMIN)", level=3)
    adm_headers = ["STT", "Mã chức năng", "Tên chức năng", "Mô tả chi tiết"]
    adm_rows = [
        ["1", "F-ADM-01", "Quản trị tài khoản & Phân quyền người dùng", "Thêm, sửa, khóa tài khoản sinh viên vi phạm; cấp tài khoản Cán bộ và gán Khoa/Phòng ban."],
        ["2", "F-ADM-02", "Quản lý danh mục Khoa/Phòng & Cấu hình SLA", "Quản lý thông tin liên hệ Khoa/Phòng; thiết lập cấu hình thời gian SLA cho từng mức ưu tiên (URGENT, MEDIUM, LOW)."],
        ["3", "F-ADM-03", "Quản trị trung tâm tri thức AI (Knowledge Hub)", "Tải lên quy chế dạng PDF, kích hoạt bóc tách Điều/Khoản, kiểm tra vector embedding và trực quan hóa vector 3D."],
        ["4", "F-ADM-04", "Quản lý ngân hàng câu hỏi FAQ chuẩn hóa", "Quản lý bộ 300+ câu hỏi chuẩn hóa và chuyển đổi các câu hỏi thực tế có lời giải xuất sắc từ Ticket thành FAQ."],
        ["5", "F-ADM-05", "Báo cáo thống kê hiệu năng & Tuân thủ SLA", "Thống kê số lượng Ticket, tỷ lệ giải quyết đúng hạn (%) của từng Khoa/Phòng, số lượng vi phạm và thời gian xử lý."],
        ["6", "F-ADM-06", "Cấu hình tích hợp Webhook an toàn", "Cấu hình Secret Token và giám sát luồng webhook nhận video render từ Microservice Node.js."]
    ]
    add_table_data(doc, adm_headers, adm_rows)

    add_heading_styled(doc, "3.2. Ma trận phân quyền tính năng & Cách ly dữ liệu theo Khoa/Phòng (RBAC Matrix)", level=2)
    rbac_headers = ["Phân hệ / Nghiệp vụ", "GUEST (Khách)", "ROLE_STUDENT", "ROLE_STAFF", "ROLE_ADMIN"]
    rbac_rows = [
        ["Đăng ký tài khoản & Xác thực Email OTP", "Không", "Có", "Không (Admin cấp)", "Không (Root cấp)"],
        ["Chatbot AI RAG tra cứu quy chế", "Có (Rate limit)", "Có (Ưu tiên)", "Có", "Có"],
        ["Tạo Ticket tư vấn (Form trực tiếp)", "Có (Nhập Email)", "Có (Tự động Profile)", "Không", "Không"],
        ["Chuyển đoạn hội thoại Chat thành Ticket", "Có", "Có", "Không", "Không"],
        ["Xem danh sách & Xử lý Ticket", "Không (chỉ qua Token)", "Không (chỉ Ticket mình)", "Có (chỉ Khoa mình)", "Có (Tất cả Khoa)"],
        ["Tiếp nhận xử lý Ticket (Claim Lock)", "Không", "Không", "Có (Gán chính chủ)", "Có (Điều phối lại)"],
        ["Đánh giá hài lòng Ticket (CSAT)", "Có (Link bảo mật)", "Có (Trên Portal)", "Không", "Không"],
        ["Xem Bảng tin chính thức & Tải tệp", "Có", "Có", "Có", "Có"],
        ["Đăng thông báo Bảng tin kèm tệp/Video", "Không", "Không", "Có", "Có"],
        ["Đăng bài Diễn đàn sinh viên", "Không", "Có (Chờ duyệt)", "Có (Duyệt ngay)", "Có (Duyệt ngay)"],
        ["Phê duyệt bài Diễn đàn", "Không", "Không", "Có (Thuộc thẩm quyền)", "Có (Toàn hệ thống)"],
        ["Thả tim (Like) & Bình luận (Comment)", "Không", "Có", "Có", "Có"],
        ["Gửi báo cáo vi phạm (Report)", "Không", "Có", "Có", "Có"],
        ["Xử lý danh sách báo cáo & Khóa bài", "Không", "Không", "Có", "Có"],
        ["Nạp tài liệu PDF & Quản trị AI Knowledge", "Không", "Không", "Không", "Có"],
        ["Quản trị người dùng & Báo cáo SLA", "Không", "Không", "Không", "Có"]
    ]
    add_table_data(doc, rbac_headers, rbac_rows)
    add_paragraph_styled(doc, "Cán bộ thuộc Khoa CNTT (department_id = 4) tuyệt đối không được phép xem hoặc can thiệp Ticket thuộc Phòng Tuyển sinh hoặc Phòng Đào tạo. Mọi truy vấn trái phép đều bị tầng bảo mật chặn đứng với lỗi 403 Forbidden.", bold_prefix="Quy tắc cách ly dữ liệu bắt buộc (Department Data Isolation): ", italic=True)

    add_heading_styled(doc, "3.3. Biểu đồ Use Case tổng quan & phân hệ", level=2)
    add_paragraph_styled(doc, "Dưới đây là các vị trí dán hình ảnh Biểu đồ Use Case được biên dịch từ các tệp mã nguồn PlantUML trong thư mục docs/uml/:")
    
    add_image_placeholder(doc, "Hình 3-1", "Biểu đồ Use Case Tổng Quan Toàn Hệ Thống QAUTE Portal", "01_usecase_overview.puml", "Biên dịch tệp docs/uml/01_usecase_overview.puml trên PlantText/VS Code rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-2", "Phân rã Use Case Phân Hệ Quản Lý Ticket & Cam Kết SLA Học Vụ", "02_usecase_ticket_sla.puml", "Biên dịch tệp docs/uml/02_usecase_ticket_sla.puml rồi dán ảnh vào đây")

    add_heading_styled(doc, "3.4. Đặc tả chi tiết các Use Case cốt lõi", level=2)
    
    # UC01
    add_heading_styled(doc, "3.4.1. Đặc tả Use Case UC01: Đăng ký tài khoản sinh viên và Xác thực Email OTP", level=3)
    uc01_headers = ["Thuộc tính", "Chi tiết đặc tả"]
    uc01_rows = [
        ["Use Case ID", "UC01"],
        ["Use Case Name", "Đăng ký tài khoản sinh viên và xác thực mã OTP qua Email trường"],
        ["Actor Chính", "Sinh viên chưa có tài khoản (GUEST chuyển đổi sang ROLE_STUDENT)"],
        ["Tiền điều kiện", "Sinh viên có hòm thư điện tử chính thức của trường (@student.hcmute.edu.vn)."],
        ["Hậu điều kiện", "Tài khoản được tạo ở trạng thái ACTIVE, được gán vai trò ROLE_STUDENT và đăng nhập được."],
        ["Luồng chính (Main Flow)", "1. Sinh viên truy cập /register, nhập Họ tên, Username, Password, Phone, Email.\n2. Nhấn nút 'Tiếp tục'.\n3. Hệ thống kiểm tra tính duy nhất của Username/Email.\n4. Hệ thống sinh mã OTP ngẫu nhiên 6 chữ số (hiệu lực 5 phút), lưu otp_tokens và gửi email bất đồng bộ.\n5. Màn hình chuyển sang giao diện nhập mã OTP.\n6. Sinh viên nhập mã OTP 6 số và nhấn 'Xác nhận'.\n7. Hệ thống xác thực OTP hợp lệ, kích hoạt tài khoản ACTIVE và đánh dấu OTP is_used = true.\n8. Hiển thị thông báo thành công và chuyển hướng đến trang đăng nhập."],
        ["Luồng rẽ nhánh", "4a. Nhấn 'Gửi lại mã OTP' sau 60 giây nếu chưa nhận được email.\n6a. Nhấn 'Hủy bỏ' để hủy quy trình."],
        ["Luồng ngoại lệ", "3a. Username hoặc Email đã tồn tại -> Hiển thị cảnh báo lỗi đỏ, dừng quy trình.\n7a. OTP sai hoặc quá hạn -> Báo lỗi không hợp lệ, tăng biến đếm nhập sai."]
    ]
    add_table_data(doc, uc01_headers, uc01_rows)

    # UC02
    add_heading_styled(doc, "3.4.2. Đặc tả Use Case UC02: Chuyển đổi cuộc hội thoại AI Chatbot thành Ticket hỗ trợ", level=3)
    uc02_headers = ["Thuộc tính", "Chi tiết đặc tả"]
    uc02_rows = [
        ["Use Case ID", "UC02"],
        ["Use Case Name", "Chuyển đổi cuộc hội thoại thành Ticket hỗ trợ chính thức & Tính toán SLA Deadline"],
        ["Actor Chính", "ROLE_STUDENT, GUEST, SYSTEM (AI Fallback Trigger)"],
        ["Tiền điều kiện", "Người dùng đang trong phiên Chat với AI Bot nhưng câu hỏi vượt quá phạm vi."],
        ["Hậu điều kiện", "Một Ticket mới được khởi tạo ở trạng thái OPEN, gán hạn chót due_date theo SLA."],
        ["Luồng chính (Main Flow)", "1. Người dùng nhấn nút 'Chuyển thành Ticket gửi Thầy/Cô' trong khung Chat.\n2. Mở Modal chuyển tiếp: Nhập Tiêu đề, chọn Khoa/Phòng, chọn Mức ưu tiên (URGENT/MEDIUM/LOW), nhập Email (nếu là Guest).\n3. Người dùng nhấn 'Gửi Ticket hỗ trợ'.\n4. Động cơ SLA tính hạn chót: URGENT (+24h), MEDIUM (+72h), LOW (+7 ngày).\n5. Sinh mã Ticket duy nhất (TK-YYYYMMDD-XXXXXX) và Token bảo mật cho Guest.\n6. Sao chép nội dung chat thành các tin nhắn trong ticket_messages.\n7. Gửi Email bất đồng bộ thông báo mã Ticket và link tra cứu.\n8. Hiển thị thông báo hoàn tất trên giao diện Chat."],
        ["Luồng ngoại lệ", "2a. Guest nhập sai định dạng Email -> Báo lỗi đỏ yêu cầu nhập đúng email."]
    ]
    add_table_data(doc, uc02_headers, uc02_rows)

    # UC03
    add_heading_styled(doc, "3.4.3. Đặc tả Use Case UC03: Cán bộ tiếp nhận (Claim) và Xử lý Ticket", level=3)
    uc03_headers = ["Thuộc tính", "Chi tiết đặc tả"]
    uc03_rows = [
        ["Use Case ID", "UC03"],
        ["Use Case Name", "Cán bộ tiếp nhận xử lý (Claim) và cập nhật tiến trình giải quyết Ticket"],
        ["Actor Chính", "ROLE_STAFF (Cán bộ phụ trách theo Khoa/Phòng)"],
        ["Tiền điều kiện", "Cán bộ đã đăng nhập thành công và có mã department_id."],
        ["Hậu điều kiện", "Trạng thái Ticket chuyển sang IN_PROGRESS (gán chính chủ) và RESOLVED khi hoàn tất."],
        ["Luồng chính (Main Flow)", "1. Cán bộ mở Dashboard Ticket thuộc Khoa/Phòng mình.\n2. Hệ thống hiển thị danh sách Ticket kèm badge thời hạn SLA.\n3. Cán bộ chọn Ticket OPEN để xem nội dung và file đính kèm.\n4. Nhấn nút 'Tiếp nhận xử lý' (Claim Ticket).\n5. Hệ thống thực thi Atomic UPDATE kiểm tra status = 'OPEN'.\n6. Ticket chuyển sang IN_PROGRESS, gán assigned_staff_id.\n7. Cán bộ nhập câu trả lời giải đáp, đính kèm văn bản hướng dẫn.\n8. Nhấn 'Hoàn thành giải quyết' (Resolve Ticket).\n9. Trạng thái chuyển sang RESOLVED, lưu resolved_at = now().\n10. Tự động gửi Email thông báo kết quả và link đánh giá CSAT cho Sinh viên."],
        ["Luồng ngoại lệ", "5a. Tranh chấp nhận Ticket (Race Condition): Cán bộ khác đã nhận trước -> rows_affected = 0, ném TicketAlreadyClaimedException và báo thông báo đã được tiếp nhận."]
    ]
    add_table_data(doc, uc03_headers, uc03_rows)

    add_heading_styled(doc, "3.5. Biểu đồ Tuần tự (Sequence Diagrams - 8 kịch bản chuẩn)", level=2)
    add_paragraph_styled(doc, "Dưới đây là các vị trí dán 8 Biểu đồ Tuần tự phản ánh chính xác luồng xử lý mã nguồn Spring Boot & Python Engine:")

    add_image_placeholder(doc, "Hình 3-3", "Sequence Diagram SD01: Tra cứu thông tin học vụ qua Trợ lý AI RAG & Fallback", "03_sd01_ai_rag_chat.puml", "Biên dịch tệp docs/uml/03_sd01_ai_rag_chat.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-4", "Sequence Diagram SD02: Chuyển đổi cuộc hội thoại thành Ticket & Thiết lập SLA", "04_sd02_convert_ticket_sla.puml", "Biên dịch tệp docs/uml/04_sd02_convert_ticket_sla.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-5", "Sequence Diagram SD03: Cán bộ Khoa tiếp nhận (Claim) và Xử lý Ticket", "05_sd03_staff_claim_ticket.puml", "Biên dịch tệp docs/uml/05_sd03_staff_claim_ticket.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-6", "Sequence Diagram SD04: Cán bộ đăng Bảng tin chính thức & Tích hợp Video Webhook Node.js", "06_sd04_official_post_video_webhook.puml", "Biên dịch tệp docs/uml/06_sd04_official_post_video_webhook.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-7", "Sequence Diagram SD05: Sinh viên đăng bài Diễn đàn & Quy trình Phê duyệt của Cán bộ", "07_sd05_student_feed_moderation.puml", "Biên dịch tệp docs/uml/07_sd05_student_feed_moderation.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-8", "Sequence Diagram SD06: Tương tác Thảo luận, Thả tim và Báo cáo vi phạm (Report)", "08_sd06_interaction_and_report.puml", "Biên dịch tệp docs/uml/08_sd06_interaction_and_report.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-9", "Sequence Diagram SD07: Quản trị viên nạp công văn PDF, tách Chunk & nhúng Vector Embedding", "09_sd07_admin_pdf_vector_ingestion.puml", "Biên dịch tệp docs/uml/09_sd07_admin_pdf_vector_ingestion.puml rồi dán ảnh vào đây")
    add_image_placeholder(doc, "Hình 3-10", "Sequence Diagram SD08: Đăng ký tài khoản sinh viên & Xác thực OTP qua Email trường", "10_sd08_auth_register_otp.puml", "Biên dịch tệp docs/uml/10_sd08_auth_register_otp.puml rồi dán ảnh vào đây")

    add_heading_styled(doc, "3.6. Biểu đồ Hoạt động (Activity Diagrams)", level=2)
    add_image_placeholder(doc, "Hình 3-11", "Activity Diagram AD01: Luồng xử lý liên thông 3 bên: Sinh viên - Chatbot AI - Cán bộ Ticket", "11_ad01_activity_process.puml", "Biên dịch tệp docs/uml/11_ad01_activity_process.puml rồi dán ảnh vào đây")

    add_heading_styled(doc, "3.7. Thiết kế Cơ sở Dữ liệu quan hệ (ERD & Data Dictionary 12 bảng)", level=2)
    add_image_placeholder(doc, "Hình 3-12", "Sơ đồ Thực Thể Liên Kết (Entity Relationship Diagram - ERD) 12 bảng chuẩn 3NF", "12_erd_database_model.puml", "Biên dịch tệp docs/uml/12_erd_database_model.puml rồi dán ảnh sơ đồ ERD vào đây")

    add_paragraph_styled(doc, "Dưới đây là từ điển dữ liệu chi tiết của 12 bảng quan hệ thực tế trong cơ sở dữ liệu MySQL 8:")

    dict_tables = [
        ("Bảng 3-9: roles (Vai trò hệ thống)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính định danh vai trò"],
            ["name", "VARCHAR(50)", "UNIQUE", "NO", "", "Mã vai trò: ROLE_STUDENT, ROLE_STAFF, ROLE_ADMIN"],
            ["description", "VARCHAR(255)", "", "YES", "NULL", "Mô tả quyền hạn của vai trò"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm tạo"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm cập nhật"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-10: departments (Khoa, Phòng ban & Đơn vị)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính định danh đơn vị"],
            ["name", "VARCHAR(150)", "UNIQUE", "NO", "", "Tên đơn vị (Phòng Đào tạo, Tuyển sinh, Khoa CNTT...)"],
            ["code", "VARCHAR(50)", "UNIQUE", "NO", "", "Mã viết tắt: DAO_TAO, TUYEN_SINH, KHOA_CNTT, DOAN_HOI"],
            ["office_location", "VARCHAR(100)", "", "YES", "NULL", "Vị trí văn phòng (Phòng A1-201, E1-402...)"],
            ["contact_email", "VARCHAR(100)", "", "YES", "NULL", "Hòm thư điện tử chính thức"],
            ["contact_phone", "VARCHAR(50)", "", "YES", "NULL", "Số điện thoại liên hệ nội bộ"],
            ["description", "TEXT", "", "YES", "NULL", "Chức năng nhiệm vụ tiếp nhận của đơn vị"],
            ["is_active", "BOOLEAN", "", "YES", "TRUE", "Trạng thái hoạt động tiếp nhận"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm khởi tạo"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm cập nhật"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-11: users (Người dùng hệ thống)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính định danh người dùng"],
            ["username", "VARCHAR(100)", "UNIQUE", "NO", "", "Tên đăng nhập (MSSV hoặc mã Cán bộ)"],
            ["password_hash", "VARCHAR(255)", "", "NO", "", "Mật khẩu mã hóa BCrypt chuẩn Spring Security"],
            ["full_name", "VARCHAR(150)", "", "NO", "", "Họ và tên đầy đủ"],
            ["email", "VARCHAR(150)", "UNIQUE", "NO", "", "Hòm thư điện tử chính danh"],
            ["phone", "VARCHAR(30)", "", "YES", "NULL", "Số điện thoại cá nhân"],
            ["avatar_url", "VARCHAR(500)", "", "YES", "NULL", "Đường dẫn ảnh đại diện"],
            ["role_id", "BIGINT", "FK", "NO", "", "Khóa ngoại liên kết bảng roles(id)"],
            ["department_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại departments(id) (bắt buộc với Cán bộ)"],
            ["status", "VARCHAR(30)", "", "NO", "'ACTIVE'", "Trạng thái: ACTIVE, PENDING_ACTIVATION, LOCKED"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm tạo tài khoản"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm cập nhật"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-12: otp_tokens (Mã xác thực Email OTP)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính"],
            ["email", "VARCHAR(150)", "", "NO", "", "Email nhận mã xác nhận OTP"],
            ["otp_code", "VARCHAR(10)", "", "NO", "", "Mã 6 số ngẫu nhiên sinh từ hệ thống"],
            ["purpose", "VARCHAR(50)", "", "NO", "", "Mục đích sử dụng: REGISTRATION, PASSWORD_RESET"],
            ["expired_at", "TIMESTAMP", "", "NO", "", "Thời điểm hết hạn hiệu lực (now + 5 phút)"],
            ["is_used", "BOOLEAN", "", "YES", "FALSE", "Đánh dấu mã đã sử dụng hay chưa"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm sinh mã"]
        ]),
        ("Bảng 3-13: tickets (Yêu cầu tư vấn học vụ & SLA)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính định danh Ticket"],
            ["ticket_code", "VARCHAR(50)", "UNIQUE", "NO", "", "Mã tra cứu duy nhất: TK-YYYYMMDD-XXXXXX"],
            ["title", "VARCHAR(255)", "", "NO", "", "Tiêu đề thắc mắc học vụ"],
            ["content", "TEXT", "", "NO", "", "Nội dung chi tiết yêu cầu giải quyết"],
            ["status", "VARCHAR(30)", "", "NO", "'OPEN'", "Trạng thái: OPEN, IN_PROGRESS, RESOLVED, CLOSED"],
            ["priority", "VARCHAR(30)", "", "NO", "'MEDIUM'", "Mức độ ưu tiên: URGENT (24h), MEDIUM (72h), LOW (7d)"],
            ["department_id", "BIGINT", "FK", "NO", "", "Khóa ngoại departments(id) chỉ định đơn vị xử lý"],
            ["creator_student_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại users(id) sinh viên tạo"],
            ["assigned_staff_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại users(id) cán bộ chịu trách nhiệm (Claim)"],
            ["guest_email", "VARCHAR(150)", "", "YES", "NULL", "Email liên hệ nếu người tạo là Khách vãng lai"],
            ["guest_name", "VARCHAR(150)", "", "YES", "NULL", "Tên người liên hệ nếu là Khách vãng lai"],
            ["access_token", "VARCHAR(100)", "", "YES", "NULL", "Chuỗi bảo mật tra cứu Ticket cho Guest"],
            ["due_date", "TIMESTAMP", "", "NO", "", "Hạn chót cam kết giải quyết theo SLA"],
            ["resolved_at", "TIMESTAMP", "", "YES", "NULL", "Thời điểm hoàn tất giải quyết"],
            ["satisfaction_rating", "INT", "", "YES", "NULL", "Điểm đánh giá mức độ hài lòng từ 1 đến 5 sao"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm mở Ticket"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm cập nhật cuối cùng"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-14: ticket_messages (Nhật ký trao đổi trong Ticket)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính tin nhắn"],
            ["ticket_id", "BIGINT", "FK", "NO", "", "Khóa ngoại tickets(id)"],
            ["sender_user_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại users(id) người gửi"],
            ["sender_type", "VARCHAR(30)", "", "NO", "", "Phân loại: STUDENT, STAFF, GUEST, SYSTEM"],
            ["message", "TEXT", "", "NO", "", "Nội dung trao đổi hoặc kết luận"],
            ["attachment_url", "VARCHAR(500)", "", "YES", "NULL", "Đường dẫn file đính kèm văn bản/minh chứng"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm gửi tin"]
        ]),
        ("Bảng 3-15: posts (Bảng tin chính thức & Diễn đàn sinh viên)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính bài viết"],
            ["title", "VARCHAR(255)", "", "NO", "", "Tiêu đề bài viết"],
            ["content", "MEDIUMTEXT", "", "NO", "", "Nội dung bài viết (hỗ trợ văn bản phong phú)"],
            ["post_type", "VARCHAR(50)", "", "NO", "", "Loại: OFFICIAL_ANNOUNCEMENT, STUDENT_DISCUSSION"],
            ["status", "VARCHAR(30)", "", "NO", "'PENDING_APPROVAL'", "Trạng thái: PENDING_APPROVAL, APPROVED, REJECTED, HIDDEN"],
            ["author_id", "BIGINT", "FK", "NO", "", "Khóa ngoại users(id) người đăng"],
            ["approved_by_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại users(id) cán bộ phê duyệt"],
            ["department_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại departments(id) đơn vị ban hành"],
            ["view_count", "INT", "", "YES", "0", "Số lượt xem bài viết"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm đăng"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm cập nhật"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-16: post_attachments (Tệp đính kèm bài viết)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính tệp đính kèm"],
            ["post_id", "BIGINT", "FK", "NO", "", "Khóa ngoại posts(id)"],
            ["file_name", "VARCHAR(255)", "", "NO", "", "Tên tệp tin gốc"],
            ["file_url", "VARCHAR(500)", "", "NO", "", "Đường dẫn lưu trữ (Cloudinary / File server)"],
            ["file_type", "VARCHAR(50)", "", "NO", "", "Định dạng: PDF, DOCX, XLSX, IMAGE_JPG, VIDEO_MP4"],
            ["file_size_bytes", "BIGINT", "", "YES", "NULL", "Dung lượng tệp tính bằng bytes"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm tải lên"]
        ]),
        ("Bảng 3-17: post_comments (Bình luận thảo luận nhiều cấp)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính bình luận"],
            ["post_id", "BIGINT", "FK", "NO", "", "Khóa ngoại posts(id)"],
            ["user_id", "BIGINT", "FK", "NO", "", "Khóa ngoại users(id) người bình luận"],
            ["parent_comment_id", "BIGINT", "FK", "YES", "NULL", "Tự liên kết bình luận cha (dạng cây)"],
            ["content", "TEXT", "", "NO", "", "Nội dung bình luận"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm gửi"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm sửa"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ]),
        ("Bảng 3-18: post_reactions (Lượt thả tim tương tác)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính"],
            ["post_id", "BIGINT", "FK", "NO", "", "Khóa ngoại posts(id)"],
            ["user_id", "BIGINT", "FK", "NO", "", "Khóa ngoại users(id)"],
            ["reaction_type", "VARCHAR(30)", "", "NO", "'LIKE'", "Loại cảm xúc: LIKE, HEART, HELPFUL"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm thả cảm xúc"]
        ]),
        ("Bảng 3-19: post_reports (Báo cáo nội dung vi phạm)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính báo cáo"],
            ["post_id", "BIGINT", "FK", "NO", "", "Khóa ngoại posts(id) bị báo cáo"],
            ["reporter_id", "BIGINT", "FK", "NO", "", "Khóa ngoại users(id) người tố cáo"],
            ["reason", "VARCHAR(255)", "", "NO", "", "Lý do: SPAM, HATE_SPEECH, MISINFORMATION..."],
            ["status", "VARCHAR(30)", "", "NO", "'PENDING'", "Trạng thái: PENDING, RESOLVED, DISMISSED"],
            ["resolved_by_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại users(id) cán bộ xử lý"],
            ["resolution_note", "VARCHAR(255)", "", "YES", "NULL", "Ghi chú biện pháp xử lý"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm gửi báo cáo"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm giải quyết"]
        ]),
        ("Bảng 3-20: ai_knowledge_documents (Kho văn bản quy chế AI RAG)", [
            ["id", "BIGINT", "PK", "NO", "AUTO_INCREMENT", "Khóa chính tài liệu"],
            ["title", "VARCHAR(255)", "", "NO", "", "Tiêu đề công văn (VD: Quy chế học vụ 2026)"],
            ["file_path", "VARCHAR(500)", "", "NO", "", "Đường dẫn lưu tệp PDF trên máy chủ"],
            ["department_id", "BIGINT", "FK", "YES", "NULL", "Khóa ngoại departments(id) cơ quan ban hành"],
            ["academic_year", "VARCHAR(20)", "", "YES", "'2026'", "Năm học áp dụng để tính toán Time-Decay"],
            ["is_active", "BOOLEAN", "", "YES", "TRUE", "Hiệu lực pháp lý của văn bản"],
            ["created_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm nạp văn bản"],
            ["updated_at", "TIMESTAMP", "", "YES", "CURRENT_TIMESTAMP", "Thời điểm sửa đổi"],
            ["is_deleted", "BOOLEAN", "", "YES", "FALSE", "Cờ xóa mềm"]
        ])
    ]
    
    dict_headers = ["Tên cột", "Kiểu dữ liệu", "Khóa", "Null", "Mặc định", "Mô tả nghiệp vụ"]
    for t_title, t_rows in dict_tables:
        add_heading_styled(doc, t_title, level=3)
        add_table_data(doc, dict_headers, t_rows)

    doc.add_page_break()

    # CHƯƠNG 4
    add_heading_styled(doc, "CHƯƠNG 4: THIẾT KẾ KIẾN TRÚC KỸ THUẬT & PHÂN HỆ AI RAG", level=1)
    
    add_heading_styled(doc, "4.1. Kiến trúc tổng thể Hybrid Dual-Engine (Spring Boot & FastAPI Python)", level=2)
    add_paragraph_styled(doc, "Hệ thống kết hợp mô hình kiến trúc lai hai tầng nhằm tận dụng tối đa thế mạnh của từng nền tảng: Spring Boot 3.3 (Cổng 8080) cho nghiệp vụ lõi, bảo mật Spring Security 6 và quản lý giao dịch MySQL 8; kết hợp với FastAPI Python (Cổng 8001) và ChromaDB cho tính toán vector embedding và đánh giá RAG Triad.")
    
    add_image_placeholder(doc, "Hình 4-1", "Sơ đồ Thành Phần Kiến Trúc Hybrid Dual-Engine (Spring Boot & FastAPI)", "13_component_architecture.puml", "Biên dịch tệp docs/uml/13_component_architecture.puml rồi dán ảnh vào đây")

    add_heading_styled(doc, "4.2. Cơ chế phân tầng tìm kiếm & Phòng thủ Quota Gemini API", level=2)
    add_paragraph_styled(doc, "Nhằm bảo vệ hệ thống không bị vượt ngưỡng hạn mức dịch vụ miễn phí (15 RPM) của Google Gemini API và tối ưu hóa độ trễ phản hồi cho sinh viên, hệ thống triển khai phễu cản tải 4 cấp độ:")
    add_paragraph_styled(doc, "Phản hồi câu hỏi chào hỏi, hỏi giờ hành chính, địa chỉ văn phòng chỉ trong < 10ms mà không tiêu tốn API.", bold_prefix="Cấp 0 (Chuẩn hóa & Regex): ")
    add_paragraph_styled(doc, "Bộ đệm ConcurrentHashMap trên RAM với chuẩn hóa Unicode NFC trả lời câu hỏi trùng lặp trong < 5ms.", bold_prefix="Cấp 1 (Local In-Memory Cache): ")
    add_paragraph_styled(doc, "Quét vector 300+ câu hỏi chuẩn hóa trong < 50ms khi độ tương đồng Cosine >= 0.85.", bold_prefix="Cấp 2 (Ngân hàng FAQ Lịch sử): ")
    add_paragraph_styled(doc, "Truy xuất Chunks công văn PDF kết hợp Gemini 2.5 Flash tổng hợp trong < 1.5s; có Deterministic Fallback khi sự cố mạng.", bold_prefix="Cấp 3 (AI RAG & Gemini Flash): ")
    
    add_image_placeholder(doc, "Hình 4-2", "Sơ đồ Phễu Cản Tải 4 Tầng & Phòng Vệ Hạn Mức Quota Gemini API", "placeholder_rag_funnel.puml", "Dán sơ đồ khối phễu cản tải 4 tầng vào đây")

    add_heading_styled(doc, "4.3. Động cơ tính hạn chót SLA & Xử lý bất đồng bộ (Async Mail / Webhook)", level=2)
    add_paragraph_styled(doc, "Tự động gán thời hạn giải quyết: URGENT (24h), MEDIUM (72h), LOW (7 ngày).", bold_prefix="1. Thuật toán cam kết thời gian SLA: ")
    add_paragraph_styled(doc, "Sử dụng câu lệnh UPDATE tickets SET assigned_staff_id = :staffId, status = 'IN_PROGRESS' WHERE id = :id AND status = 'OPEN' để triệt tiêu tranh chấp nhận Ticket.", bold_prefix="2. Thuật toán Atomic Claim chống Race Condition: ")
    add_paragraph_styled(doc, "Xác thực chữ ký HMAC-SHA256 trên Header X-Signature-SHA256 trước khi lưu tệp Video MP4 từ Microservice Node.js.", bold_prefix="3. Bảo mật Webhook tích hợp Video: ")

    output_path = "BaoCao_KhaoSat_Va_ThietKe_QAUTE_Portal.docx"
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    build_word_document()
